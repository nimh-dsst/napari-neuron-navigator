"""CCF policy UI isolation and durable provenance regressions."""

import json
import os
import subprocess
import sys
import textwrap
from dataclasses import replace

import pytest
from openpyxl import load_workbook

from napari_neuron_navigator.analysis.pearson_policy import (
    ccf_pearson_mode_text,
    ccf_pearson_policy_metadata,
)


@pytest.mark.parametrize("corrected", [True, False, None])
def test_policy_provenance_round_trips_exports_and_project(tmp_path, corrected):
    from napari_neuron_navigator.analysis.export import (
        export_cluster_workbook,
        export_distance_workbook,
    )
    from napari_neuron_navigator.cluster_assignments import ClusterAssignmentStore
    from napari_neuron_navigator.project_io import (
        load_project_bundle,
        save_project_bundle,
    )
    from tests.test_analysis_exports import _make_cluster_result, _write_source_parquet

    source = tmp_path / "source.parquet"
    _write_source_parquet(source)
    result, colors = _make_cluster_result(source)
    extra = (
        {}
        if corrected is None
        else {"correlation": ccf_pearson_policy_metadata(corrected)}
    )
    result.metadata = replace(result.metadata, extra_metadata=extra)
    for export in (export_cluster_workbook, export_distance_workbook):
        workbook = tmp_path / f"{export.__name__}.xlsx"
        export(workbook, result, colors)
        loaded = load_workbook(workbook)
        metadata = dict(loaded["Metadata"].iter_rows(min_row=2, values_only=True))
        assert json.loads(metadata["extra_metadata"]) == extra
        loaded.close()

    store = ClusterAssignmentStore()
    store.add_result(
        result, method_name="Voxel Correlation", run_metadata=result.metadata.to_dict()
    )
    table_state = {
        "version": 2,
        "entries": [{"file_id": file_id} for file_id in result.neuron_ids],
        "cluster_assignments": store.to_state(),
    }
    bundle_path = tmp_path / "pearson.nnproj"
    save_project_bundle(
        bundle_path, source_parquet_path=source, table_state=table_state
    )
    bundle = load_project_bundle(bundle_path)
    restored = ClusterAssignmentStore.from_state(
        bundle.table_state["cluster_assignments"]
    )
    assert restored.active.assignments == store.active.assignments
    assert restored.active.run_metadata["extra_metadata"] == extra
    assert restored.active.runtime_result is None
    if corrected is None:
        assert ccf_pearson_mode_text(extra) == ""


def test_real_checkboxes_default_visibility_independence_and_busy_states(tmp_path):
    """Use fresh Qt imports to avoid the suite's intentionally stubbed modules."""
    script = textwrap.dedent("""
        from types import SimpleNamespace
        from qtpy.QtWidgets import QApplication
        from napari_neuron_navigator.widgets.analysis_tab import AnalysisTabWidget
        from napari_neuron_navigator.widgets.search_tab import SearchTabWidget
        from tests.test_analysis_region_options import _DummyViewer

        app = QApplication.instance() or QApplication([])
        analysis = AnalysisTabWidget(_DummyViewer())
        search = SearchTabWidget()
        analysis.refresh_flatmap_coordinate_availability = lambda: None
        analysis._clustering_method_combo.setCurrentText("Voxel Correlation")
        for widget in (analysis, search):
            checkbox = widget._corrected_pearson_cb
            assert checkbox.text() == "Use corrected Pearson correlation"
            assert checkbox.isChecked()
            assert not checkbox.isHidden()
            assert "zero-variance" in checkbox.toolTip()

        analysis._corrected_pearson_cb.setChecked(False)
        assert search._corrected_pearson_cb.isChecked()
        analysis._clustering_method_combo.setCurrentText("Soma Location")
        assert analysis._corrected_pearson_cb.isHidden()
        analysis._clustering_method_combo.setCurrentText("Voxel Correlation")
        assert not analysis._corrected_pearson_cb.isHidden()
        assert not analysis._corrected_pearson_cb.isChecked()

        analysis._coordinate_space_combo.addItem("Flat map + Depth")
        analysis._coordinate_space_combo.setCurrentText("Flat map + Depth")
        assert analysis._corrected_pearson_cb.isHidden()
        analysis._coordinate_space_combo.setCurrentText("CCFv3 Coordinates")
        assert not analysis._corrected_pearson_cb.isChecked()
        assert not analysis._corrected_pearson_cb.isHidden()

        search._corrected_pearson_cb.setChecked(False)
        search._coordinate_space_combo.addItem("Flat map + Depth", "flatmap")
        search._coordinate_space_combo.setCurrentIndex(1)
        assert search._corrected_pearson_cb.isHidden()
        search._coordinate_space_combo.setCurrentIndex(0)
        assert not search._corrected_pearson_cb.isHidden()
        assert not search._corrected_pearson_cb.isChecked()

        for widget in (analysis, search):
            # Controls must lock even before the thread starts running.
            widget._worker_thread = SimpleNamespace(isRunning=lambda: False)
            widget._update_button_states()
            assert not widget._corrected_pearson_cb.isEnabled()
            widget._worker_thread = None
        analysis._pending_clustering_request = object()
        search._pending_request = object()
        for widget in (analysis, search):
            widget._update_button_states()
            assert not widget._corrected_pearson_cb.isEnabled()
        analysis._pending_clustering_request = None
        search._pending_request = None
        for widget in (analysis, search):
            widget._update_button_states()
            assert widget._corrected_pearson_cb.isEnabled()
            assert not widget._corrected_pearson_cb.isChecked()
            widget.close()
    """)
    completed = subprocess.run(
        [sys.executable, "-c", script],
        env={
            **os.environ,
            "QT_QPA_PLATFORM": "offscreen",
            "MPLCONFIGDIR": str(tmp_path),
        },
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert completed.returncode == 0, completed.stdout + completed.stderr
